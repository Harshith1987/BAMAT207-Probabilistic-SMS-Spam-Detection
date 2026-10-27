import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report_docx(output_path, plots_dir):
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("BAMAT207: APPLICATIONS OF PROBABILITY AND STATISTICS")
    title_run.font.name = "Arial"
    title_run.font.size = Pt(16)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 78, 121)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Beyond the Words: A Probabilistic & Bayesian Approach to SMS Spam Detection")
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(13)
    sub_run.font.bold = True
    sub_run.font.color.rgb = RGBColor(89, 89, 89)
    
    doc.add_paragraph()
    
    # Metadata Box
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course Name & Code:", "Applications of Probability and Statistics (BAMAT207)"),
        ("Topic:", "Probability and Random Variables"),
        ("Student Name / ID:", "Harshith (Harshith1987)"),
        ("Submission Date:", "27.10.2026"),
        ("GitHub Repository Link:", "https://github.com/Harshith1987/BAMAT207-Probabilistic-SMS-Spam-Detection")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.text = k
        c1.paragraphs[0].runs[0].font.bold = True
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c2.text = v
        c2.paragraphs[0].runs[0].font.size = Pt(10)
        set_cell_background(c1, "F2F2F2")
        set_cell_background(c2, "FAFAFA")
        
    doc.add_paragraph()
    
    # 1. Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Problem Formulation", level=1)
    h1.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    p = doc.add_paragraph(
        "Unsolicited mobile text messages (SMS spam) present critical challenges in digital communications, "
        "including fraud, phishing, and malware distribution. The objective of this project is to develop a "
        "mathematically rigorous probabilistic classification system that categorizes an incoming SMS message as either "
        "Spam (Y = 1) or Genuine / Ham (Y = 0).\n\n"
        "The project applies foundational principles of Probability and Random Variables: computing prior probabilities, "
        "evaluating word-level empirical conditional probabilities, mathematically proving and verifying Bayes' Theorem, "
        "formulating the classification process as a discrete Bernoulli Random Variable (calculating Expectation and Variance), "
        "training a Multinomial Naive Bayes classifier with Laplace smoothing, and conducting an in-depth confusion matrix "
        "error audit of False Positives and False Negatives."
    )
    
    # 2. Dataset Handling & Class Distribution
    h2 = doc.add_heading("2. Dataset Handling & Prior Probabilities", level=1)
    h2.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    doc.add_paragraph(
        "The project uses the SMS Spam Collection dataset from the UCI Machine Learning Repository (ID: 228). "
        "Data hygiene was enforced by validating null values and removing duplicate entries to prevent estimation bias."
    )
    
    # Dataset Table
    ds_table = doc.add_table(rows=6, cols=3)
    ds_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Metric / Parameter", "Value", "Statistical Interpretation"]
    for j, h in enumerate(headers):
        cell = ds_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        
    ds_data = [
        ("Raw Messages", "5,572", "Original corpus from UCI repository"),
        ("Missing / Null Values", "0", "No data imputation required"),
        ("Duplicate Records Removed", "403", "Eliminated redundant messages"),
        ("Clean Unique Messages (N)", "5,169", "Sample space size for probability analysis"),
        ("Prior Probability P(Ham)", "0.8737 (4,516 msgs)", "87.37% baseline probability of genuine SMS"),
        ("Prior Probability P(Spam)", "0.1263 (653 msgs)", "12.63% baseline probability of spam SMS")
    ]
    for i, (m, val, desc) in enumerate(ds_data, start=1):
        if i < len(ds_table.rows):
            row = ds_table.rows[i]
            row.cells[0].text = m
            row.cells[1].text = val
            row.cells[2].text = desc
            for c in row.cells:
                c.paragraphs[0].runs[0].font.size = Pt(9.5)
                if i % 2 == 1:
                    set_cell_background(c, "F9FBFD")
                    
    doc.add_paragraph()
    
    # Class Distribution Image
    class_plot_path = os.path.join(plots_dir, "class_distribution.png")
    if os.path.exists(class_plot_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(class_plot_path, width=Inches(5.0))
        caption = doc.add_paragraph("Figure 1: SMS Dataset Class Distribution (Ham vs Spam)")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].font.size = Pt(9)
        caption.runs[0].font.italic = True
        
    # 3. Word-Based Conditional Probability & Bayes' Theorem
    h3 = doc.add_heading("3. Word-Based Probability Analysis & Bayes' Theorem Verification", level=1)
    h3.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    doc.add_paragraph(
        "For any suspicious word W, we calculate the conditional probability P(Spam | W) via two independent methods:\n"
        "1. Empirical Ratio: P(Spam | W) = Count(Spam containing W) / Count(All messages containing W)\n"
        "2. Bayes' Theorem: P(Spam | W) = [P(W | Spam) * P(Spam)] / P(W)\n"
        "Both formulations produce identical numerical results, demonstrating the theorem in practice."
    )
    
    # Word Table
    word_table = doc.add_table(rows=6, cols=6)
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_headers = ["Word (W)", "Total (N_W)", "Spam (N_S)", "P(W|Spam)", "P(Spam|W) Bayes", "Equivalence Verified"]
    for j, h in enumerate(w_headers):
        cell = word_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        
    word_data = [
        ("free", "203", "148", "0.2266", "0.7291 (72.91%)", "True"),
        ("prize", "73", "73", "0.1118", "1.0000 (100.0%)", "True"),
        ("win", "55", "49", "0.0750", "0.8909 (89.09%)", "True"),
        ("urgent", "62", "57", "0.0873", "0.9194 (91.94%)", "True"),
        ("offer", "30", "24", "0.0368", "0.8000 (80.00%)", "True")
    ]
    for i, row_data in enumerate(word_data, start=1):
        row = word_table.rows[i]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
            row.cells[j].paragraphs[0].runs[0].font.size = Pt(9)
            if i % 2 == 1:
                set_cell_background(row.cells[j], "F9FBFD")
                
    doc.add_paragraph()
    
    # Word plot
    word_plot_path = os.path.join(plots_dir, "word_probabilities.png")
    if os.path.exists(word_plot_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(word_plot_path, width=Inches(5.0))
        caption = doc.add_paragraph("Figure 2: Conditional Probability of Spam Given Suspicious Words: P(Spam | Word)")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].font.size = Pt(9)
        caption.runs[0].font.italic = True

    # 4. Random Variable Analysis
    h4 = doc.add_heading("4. Random Variable Analysis (Bernoulli Distribution)", level=1)
    h4.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    doc.add_paragraph(
        "Let the random variable Y define the classification state of an arbitrary message:\n"
        "  Y = 1 if the message is Spam\n"
        "  Y = 0 if the message is Ham\n\n"
        "Here, Y follows a Bernoulli distribution: Y ~ Bernoulli(p) where p = P(Spam) = 0.1263 and q = 1 - p = 0.8737.\n\n"
        "• Expectation E(Y):\n"
        "  E(Y) = sum(y * P(Y = y)) = (1 * p) + (0 * q) = p = 0.1263\n\n"
        "• Variance Var(Y):\n"
        "  Var(Y) = E[Y^2] - (E[Y])^2 = p - p^2 = p(1 - p) = 0.1263 * 0.8737 = 0.1104\n\n"
        "• Standard Deviation SD(Y):\n"
        "  SD(Y) = sqrt(Var(Y)) = sqrt(0.1104) = 0.3322"
    )

    # 5. Data Preparation & Model Training
    h5 = doc.add_heading("5. Model Building: Multinomial Naive Bayes", level=1)
    h5.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    doc.add_paragraph(
        "• Data Partitioning: 80% training set (4,135 samples) and 20% testing set (1,034 samples) "
        "using stratified sampling (stratify=y) to maintain identical class proportions.\n"
        "• Feature Extraction: CountVectorizer extracted 7,414 unique word tokens as features.\n"
        "• Classification Algorithm: Multinomial Naive Bayes (MultinomialNB) with Laplace Add-1 smoothing (alpha = 1.0)."
    )

    # 6. Evaluation & Confusion Matrix
    h6 = doc.add_heading("6. Model Evaluation & Error Analysis", level=1)
    h6.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    eval_table = doc.add_table(rows=5, cols=2)
    eval_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    eval_metrics = [
        ("Overall Test Accuracy", "98.36% (0.9836)"),
        ("Precision (Spam)", "95.97% (0.9597)"),
        ("Recall (Spam)", "90.84% (0.9084)"),
        ("F1-Score (Spam)", "0.9333")
    ]
    eval_table.rows[0].cells[0].text = "Evaluation Metric"
    eval_table.rows[0].cells[1].text = "Model Score"
    eval_table.rows[0].cells[0].paragraphs[0].runs[0].font.bold = True
    eval_table.rows[0].cells[1].paragraphs[0].runs[0].font.bold = True
    set_cell_background(eval_table.rows[0].cells[0], "1F4E79")
    set_cell_background(eval_table.rows[0].cells[1], "1F4E79")
    eval_table.rows[0].cells[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    eval_table.rows[0].cells[1].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    
    for i, (m, score) in enumerate(eval_metrics, start=1):
        row = eval_table.rows[i]
        row.cells[0].text = m
        row.cells[1].text = score
        row.cells[0].paragraphs[0].runs[0].font.size = Pt(10)
        row.cells[1].paragraphs[0].runs[0].font.size = Pt(10)
        
    doc.add_paragraph()
    
    # Confusion matrix plot
    cm_plot_path = os.path.join(plots_dir, "confusion_matrix.png")
    if os.path.exists(cm_plot_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(cm_plot_path, width=Inches(4.5))
        caption = doc.add_paragraph("Figure 3: Confusion Matrix (Multinomial Naive Bayes on 1,034 Test Samples)")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].font.size = Pt(9)
        caption.runs[0].font.italic = True
        
    doc.add_paragraph(
        "Critical Error Audit:\n"
        "1. False Positives (Genuine messages classified as Spam): 5 messages (0.55% error rate on Ham). "
        "High precision is paramount to ensure legitimate personal or business messages are not filtered out.\n"
        "2. False Negatives (Spam messages classified as Genuine): 12 messages."
    )

    # 7. Real-time Testing
    h7 = doc.add_heading("7. Real-Time Testing & Inference Demonstrations", level=1)
    h7.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    test_table = doc.add_table(rows=6, cols=4)
    test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_headers = ["Input SMS Message", "Prediction", "P(Spam|Msg)", "P(Ham|Msg)"]
    for j, h in enumerate(t_headers):
        cell = test_table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(9)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1F4E79")
        
    test_data = [
        ("Congratulations! You have won a free prize. Claim now.", "Spam", "1.0000 (100.0%)", "0.0000 (0.0%)"),
        ("Hey, are you free this afternoon to study for the BAMAT207 exam?", "Ham", "0.0006 (0.06%)", "0.9994 (99.94%)"),
        ("URGENT! Your mobile account has won an exclusive $1,000 cash offer.", "Spam", "1.0000 (100.0%)", "0.0000 (0.0%)"),
        ("Can you send me the lecture notes when you get home?", "Ham", "0.0028 (0.28%)", "0.9972 (99.72%)"),
        ("Win a brand new car today! Free entry, text WIN to 80085 immediately.", "Spam", "0.9999 (99.99%)", "0.0001 (0.01%)")
    ]
    for i, row_data in enumerate(test_data, start=1):
        row = test_table.rows[i]
        for j, val in enumerate(row_data):
            row.cells[j].text = val
            row.cells[j].paragraphs[0].runs[0].font.size = Pt(8.5)
            if i % 2 == 1:
                set_cell_background(row.cells[j], "F9FBFD")
                
    doc.add_paragraph()

    # 8. Conclusion
    h8 = doc.add_heading("8. Conclusion & Learning Outcomes", level=1)
    h8.runs[0].font.color.rgb = RGBColor(31, 78, 121)
    
    doc.add_paragraph(
        "1. Integrated foundational probability theory, conditional probability, Bayes' Theorem, and random variables into a functional AI classifier.\n"
        "2. Demonstrated that the Bernoulli random variable variance is driven by class balance, and verified Bayes' theorem both analytically and empirically.\n"
        "3. Evaluated a complete NLP pipeline achieving 98.36% accuracy and 95.97% precision, establishing practical trade-offs between precision and recall.\n"
        "4. Established a professional GitHub repository profile for coursework submission: https://github.com/Harshith1987/BAMAT207-Probabilistic-SMS-Spam-Detection"
    )
    
    doc.save(output_path)
    print(f"Report document successfully created at: {output_path}")

if __name__ == "__main__":
    base_dir = r"C:\Users\acs\.gemini\antigravity\scratch\sms_spam_detection"
    plots_dir = os.path.join(base_dir, "plots")
    out_file = os.path.join(base_dir, "BAMAT207_Probabilistic_SMS_Spam_Detection_Report.docx")
    create_report_docx(out_file, plots_dir)
